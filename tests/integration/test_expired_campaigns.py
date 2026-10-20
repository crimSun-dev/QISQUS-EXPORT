"""Qiscus deletes old campaigns' recipient logs (~14 months) and answers 400 'broadcast_job not found'.
Those are recorded once as 'not kept', never retried, and never stop the run or make it partial."""

import httpx
import respx
from openpyxl import load_workbook

from qexport.orchestrator import EXIT_OK, EXIT_PARTIAL
from tests.conftest import comment, comments_page, jobs_page, logs_page, raw_job, raw_log, raw_room, rooms_page
from tests.integration.test_pipeline import COMMENTS_URL, JOBS_URL, ROOMS_URL, build, logs_url

NOT_FOUND = {"errors": {"message": "broadcast_job not found"}, "status": 400}


def _tenant_with_expired(n_old=20):
    respx.post(ROOMS_URL).mock(return_value=httpx.Response(200, json=rooms_page([raw_room(1)], None)))
    respx.get(COMMENTS_URL).mock(return_value=httpx.Response(200, json=comments_page([
        comment(1, 491000001, email="6281200000001", text="Halo")])))
    jobs = [raw_job(900, "Recent", created="2026-10-01T08:00:00")]
    jobs += [raw_job(800 + i, f"Old {i}", created=f"2025-07-{10 + i % 18:02d}T08:00:00") for i in range(n_old)]
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page(jobs, None)))
    respx.get(logs_url(900)).mock(return_value=httpx.Response(200, json=logs_page(
        [raw_log(1, "6281200000001")], None)))
    return {800 + i: respx.get(logs_url(800 + i)).mock(return_value=httpx.Response(400, json=NOT_FOUND))
            for i in range(n_old)}


@respx.mock
def test_expired_campaigns_do_not_stop_the_run_or_make_it_partial(settings, client, store, norm):
    old = _tenant_with_expired(20)                     # more than the 15-in-a-row circuit breaker
    result = build(settings, client, store, norm).run()
    assert result.exit_code == EXIT_OK, result.message
    assert all(r.call_count == 1 for r in old.values())
    assert store.skipped() == [] and store.jobs_gone()[0] == 20
    assert result.report.complete and not result.report.path.stem.endswith("_PARTIAL")

    wb = load_workbook(result.report.path, read_only=True)
    info = {r[0]: r[1] for r in wb["Export Info"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert info["Campaigns Qiscus no longer keeps recipients for"].startswith("20 (created up to 2025-07-")
    camp = {r[0]: r[19] for r in wb["Campaigns"].iter_rows(min_row=2, values_only=True)}
    assert camp["Old 0"] == "not kept by Qiscus" and camp["Recent"] == "done"


@respx.mock
def test_expired_campaigns_are_never_retried(settings, client, store, norm):
    old = _tenant_with_expired(3)
    build(settings, client, store, norm).run(write_xlsx=False)
    assert store.reset_failed() == 0                   # "Retry failed items" leaves them alone
    build(settings, client, store, norm).run(write_xlsx=False)
    assert all(r.call_count == 1 for r in old.values())


@respx.mock
def test_other_400s_still_trip_the_circuit_breaker(settings, client, store, norm):
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page(
        [raw_job(1000 + n, f"Job {n}") for n in range(20)], None)))
    for n in range(20):
        respx.get(logs_url(1000 + n)).mock(return_value=httpx.Response(400, json={"errors": "queued"}))
    result = build(settings, client, store, norm).run(write_xlsx=False, include_messages=False)
    assert result.exit_code == EXIT_PARTIAL and "stopped" in result.message and store.jobs_gone()[0] == 0


@respx.mock
def test_conversations_run_is_not_partial_because_of_broadcast_failures(settings, client, store, norm):
    respx.post(ROOMS_URL).mock(return_value=httpx.Response(200, json=rooms_page([raw_room(1)], None)))
    respx.get(COMMENTS_URL).mock(return_value=httpx.Response(200, json=comments_page([])))
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page([raw_job(5, "J")], None)))
    respx.get(logs_url(5)).mock(return_value=httpx.Response(400, json={"errors": "queued"}))
    build(settings, client, store, norm).run(write_xlsx=False, include_messages=False)
    assert store.skipped()                             # a broadcast job failed
    result = build(settings, client, store, norm, broadcasts=False).run(write_xlsx=False)
    assert result.exit_code == EXIT_OK
