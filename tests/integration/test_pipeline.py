import json
from pathlib import Path

import httpx
import pytest
import respx
from openpyxl import load_workbook

from qexport.api.models import RawRoom
from qexport.extract.broadcasts import BroadcastEnumerator, BroadcastLogFetcher
from qexport.extract.messages import CursorLoopError, MessageFetcher
from qexport.extract.rooms import RoomEnumerator
from qexport.load import excel as excel_mod
from qexport.load.excel import (
    ACTIONS_COLUMNS, CAMPAIGNS_COLUMNS, CONVERSATIONS_COLUMNS, LEADS_COLUMNS, MESSAGES_COLUMNS,
    RECIPIENTS_COLUMNS, ExcelWriter,
)
from qexport.orchestrator import EXIT_AUTH, EXIT_OK, EXIT_PARTIAL, ExportOrchestrator
from qexport.store.checkpoint import CheckpointStore, StoreAppMismatchError
from tests.conftest import (
    ADMIN_EMAIL, APP, OMNI, SDK, comment, comments_page, jobs_page, logs_page, raw_job, raw_log, raw_room,
    rooms_page,
)

ROOMS_URL = f"{OMNI}/api/v2/customer_rooms"
COMMENTS_URL = f"{SDK}/api/v2/sdk/load_comments"
JOBS_URL = f"{OMNI}/api/v3/admin/broadcast_jobs"


def logs_url(job_id):
    return f"{OMNI}/api/v2/admin/broadcast_logs/{job_id}"


def build(settings, client, store, norm, *, broadcasts=True):
    return ExportOrchestrator(
        settings, client, store, RoomEnumerator(client, store, norm), MessageFetcher(client, norm),
        ExcelWriter(norm), None,
        BroadcastEnumerator(client, store, norm) if broadcasts else None,
        BroadcastLogFetcher(client, norm) if broadcasts else None)


def rooms_by_cursor(pages: dict):
    """respx side effect: pick the page by the request's cursor_after (None = first page)."""
    def handler(request):
        cur = json.loads(request.content).get("cursor_after")
        return httpx.Response(200, json=pages[cur])
    return handler


def comments_by_room(per_room: dict):
    def handler(request):
        rid = request.url.params["topic_id"]
        if rid in per_room and isinstance(per_room[rid], int):
            return httpx.Response(per_room[rid])
        return httpx.Response(200, json=comments_page(per_room.get(rid, [])))
    return handler


# --------------------------------------------------------------------------- store

def test_store_rejects_db_created_for_different_app_code(tmp_path):
    CheckpointStore(tmp_path / "s.sqlite3", "app-a").close()
    with pytest.raises(StoreAppMismatchError):
        CheckpointStore(tmp_path / "s.sqlite3", "app-b")


def test_upsert_rooms_page_advances_cursor_atomically(store, norm, monkeypatch):
    rooms = [norm.room(RawRoom.model_validate(raw_room(i))) for i in range(3)]
    store.upsert_rooms_page(rooms[:1], "c1")
    bad = rooms[1:] + ["not a room"]   # blows up mid-transaction
    with pytest.raises(AttributeError):
        store.upsert_rooms_page(bad, "c2")   # type: ignore[arg-type]
    assert store.get_room_cursor() == "c1" and store.progress().rooms_total == 1


def test_commit_room_replaces_messages_and_marks_done_in_one_transaction(store, norm):
    room = norm.room(RawRoom.model_validate(raw_room(1)))
    store.upsert_rooms_page([room], None)
    from qexport.api.models import RawComment

    m1 = [norm.message(RawComment.model_validate(comment(i, 491000001, email="x")), room) for i in (1, 2)]
    store.commit_room(room.room_id, m1, None)
    store.commit_room(room.room_id, m1[:1], None)
    p = store.progress()
    assert p.messages == 1 and p.done == 1


def test_room_failing_max_attempts_becomes_skipped(store, norm):
    room = norm.room(RawRoom.model_validate(raw_room(1)))
    store.upsert_rooms_page([room], None)
    for _ in range(4):
        store.mark_room_failed(room.room_id, "boom", 5)
    assert store.progress().failed == 1
    store.mark_room_failed(room.room_id, "boom", 5)
    assert store.progress().skipped == 1 and list(store.iter_rooms_to_fetch(5)) == []


# --------------------------------------------------------------------------- enumerators / fetchers

@respx.mock
def test_enumerator_resumes_from_saved_cursor(store, client, norm):
    store.upsert_rooms_page([], "seeded-cursor")
    route = respx.post(ROOMS_URL).mock(return_value=httpx.Response(200, json=rooms_page([], None)))
    RoomEnumerator(client, store, norm).run()
    assert json.loads(route.calls[0].request.content)["cursor_after"] == "seeded-cursor"


@respx.mock
def test_enumerator_stops_when_cursor_repeats(store, client, norm):
    route = respx.post(ROOMS_URL).mock(side_effect=rooms_by_cursor({
        None: rooms_page([raw_room(1)], "A"), "A": rooms_page([raw_room(2)], "A")}))
    assert RoomEnumerator(client, store, norm).run() == 2
    assert route.call_count == 2 and store.enumeration_done()


@respx.mock
def test_enumerator_stops_on_empty_page(store, client, norm):
    respx.post(ROOMS_URL).mock(side_effect=rooms_by_cursor({
        None: rooms_page([raw_room(1)], "A"), "A": rooms_page([], "B")}))
    assert RoomEnumerator(client, store, norm).run() == 1


@respx.mock
def test_enumerator_deduplicates_room_seen_on_two_pages(store, client, norm):
    respx.post(ROOMS_URL).mock(side_effect=rooms_by_cursor({
        None: rooms_page([raw_room(1), raw_room(2)], "A"), "A": rooms_page([raw_room(2), raw_room(3)], None)}))
    assert RoomEnumerator(client, store, norm).run() == 3


@respx.mock
def test_fetcher_pages_backwards_until_short_page(client, norm):
    room = norm.room(RawRoom.model_validate(raw_room(1)))
    pages = {None: range(337, 237, -1), "238": range(237, 137, -1), "138": range(137, 100, -1)}

    def handler(request):
        key = request.url.params.get("last_comment_id")
        return httpx.Response(200, json=comments_page(
            [comment(i, 491000001, email="x", minutes=i) for i in pages[key]]))
    route = respx.get(COMMENTS_URL).mock(side_effect=handler)
    msgs = MessageFetcher(client, norm).fetch_room(room)
    assert len(msgs) == 237 and route.call_count == 3
    assert [m.msg_id for m in msgs] == sorted(m.msg_id for m in msgs)


@respx.mock
def test_fetcher_raises_cursor_loop_error_on_repeated_oldest_id(client, norm):
    room = norm.room(RawRoom.model_validate(raw_room(1)))
    respx.get(COMMENTS_URL).mock(return_value=httpx.Response(200, json=comments_page(
        [comment(i, 491000001, email="x") for i in range(100, 0, -1)])))
    with pytest.raises(CursorLoopError):
        MessageFetcher(client, norm).fetch_room(room)


@respx.mock
def test_broadcast_log_fetcher_follows_cursor(client, norm):
    def handler(request):
        cur = request.url.params.get("cursor_after")
        if cur is None:
            return httpx.Response(200, json=logs_page([raw_log(i, f"62812{i:08d}") for i in range(100)], 99))
        return httpx.Response(200, json=logs_page([raw_log(i, f"62812{i:08d}") for i in range(100, 130)], None))
    respx.get(logs_url(7)).mock(side_effect=handler)
    assert len(BroadcastLogFetcher(client, norm).fetch_job(7)) == 130


# --------------------------------------------------------------------------- orchestrator

def _tenant(n_rooms=5, forbidden=(), broken_auth=False):
    """Mocks a small tenant: rooms 1..n, one broadcast job that reached rooms 1 and 2;
    room 1's customer tapped the button, room 2 asked a question."""
    respx.post(ROOMS_URL).mock(side_effect=rooms_by_cursor({
        None: rooms_page([raw_room(i) for i in range(1, n_rooms + 1)], None)}))
    per_room: dict = {}
    for i in range(1, n_rooms + 1):
        rid = 491000000 + i
        phone = f"62812000{i:05d}"
        per_room[str(rid)] = [
            comment(rid * 10 + 1, rid, email=ADMIN_EMAIL, text="Hi! Build My University Plan", minutes=0,
                    unique_temp_id=f"wa_wamid.R{i}", extras={"broadcast_message": {"name": "welcome_by_qya2"}}),
        ]
        if i == 1:
            per_room[str(rid)].append(comment(
                rid * 10 + 2, rid, email=phone, type_="reply", text="Build My University Plan", minutes=5,
                payload={"replied_comment_id": rid * 10 + 1, "replied_comment_message": "Hi!"},
                extras={"sub_type": "button", "button": {"text": "Build My University Plan"},
                        "context": {"id": "wamid.R1"}}))
        if i == 2:
            per_room[str(rid)].append(comment(rid * 10 + 2, rid, email=phone, text="Berapa biayanya?",
                                              minutes=30))
    for i in forbidden:
        per_room[str(491000000 + i)] = 403
    if broken_auth:
        per_room[str(491000003)] = 401
    respx.get(COMMENTS_URL).mock(side_effect=comments_by_room(per_room))
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page(
        [raw_job(900, "Qya_First_BM", recipients=3)], None)))
    respx.get(logs_url(900)).mock(return_value=httpx.Response(200, json=logs_page([
        raw_log(1, "6281200000001", status=4, wamid="wamid.R1"),
        raw_log(2, "6281200000002", status=3, wamid="wamid.R2"),
        raw_log(3, "6281299999999", status=-1, notes="Error 131026, Message undeliverable"),
    ], None)))


@respx.mock
def test_room_access_error_marks_failed_and_run_continues(settings, client, store, norm):
    _tenant(forbidden=(3,))
    result = build(settings, client, store, norm).run(write_xlsx=False)
    p = store.progress()
    assert (p.done, p.failed) == (4, 1) and result.exit_code == EXIT_PARTIAL


@respx.mock
def test_auth_error_mid_fetch_exits_with_code_2_and_state_resumable(settings, client, store, norm):
    _tenant(broken_auth=True)
    result = build(settings, client, store, norm).run(write_xlsx=False)
    assert result.exit_code == EXIT_AUTH
    done_before = store.progress().done
    assert 0 < done_before < 5
    respx.reset()
    _tenant()
    result = build(settings, client, store, norm).run(write_xlsx=False)
    assert result.exit_code == EXIT_OK and store.progress().done == 5


@respx.mock
def test_keyboard_interrupt_leaves_room_pending_and_no_partial_messages(settings, client, store, norm,
                                                                        monkeypatch):
    _tenant()
    orch = build(settings, client, store, norm)
    real = MessageFetcher.fetch_room
    calls = {"n": 0}

    def flaky(self, room):
        calls["n"] += 1
        if calls["n"] == 3:
            raise KeyboardInterrupt
        return real(self, room)
    monkeypatch.setattr(MessageFetcher, "fetch_room", flaky)
    result = orch.run(write_xlsx=False)
    assert result.exit_code == 130
    p = store.progress()
    assert p.done == 2 and p.pending == 3


@respx.mock
def test_since_last_run_refetches_only_rooms_with_new_activity(settings, client, store, norm):
    _tenant()
    build(settings, client, store, norm).run(write_xlsx=False)
    respx.reset()
    _tenant()
    bumped = [raw_room(i) for i in range(1, 6)]
    bumped[0]["last_comment_timestamp"] = "2026-10-20T00:00:00Z"
    bumped[1]["last_comment_timestamp"] = "2026-10-20T00:01:00Z"
    respx.post(ROOMS_URL).mock(side_effect=rooms_by_cursor({None: rooms_page(bumped, None)}))
    route = respx.get(COMMENTS_URL).mock(return_value=httpx.Response(200, json=comments_page([])))
    build(settings, client, store, norm).run(since_last_run=True, write_xlsx=False)
    assert route.call_count == 2


# --------------------------------------------------------------------------- excel

@respx.mock
def test_full_export_workbook_contents(settings, client, store, norm, tmp_path):
    _tenant()
    out = tmp_path / "out.xlsx"
    result = build(settings, client, store, norm).run(out_path=out)
    assert result.exit_code == EXIT_OK
    assert not out.with_name(out.name + ".partial").exists()
    wb = load_workbook(out, read_only=True)
    assert wb.sheetnames == ["Leads", "Actions", "Campaigns", "Broadcast Recipients", "Conversations",
                             "Messages", "Export Info"]
    expected = {"Leads": LEADS_COLUMNS, "Actions": ACTIONS_COLUMNS, "Campaigns": CAMPAIGNS_COLUMNS,
                "Broadcast Recipients": RECIPIENTS_COLUMNS, "Conversations": CONVERSATIONS_COLUMNS,
                "Messages": MESSAGES_COLUMNS}
    for name, cols in expected.items():
        assert tuple(c.value for c in next(wb[name].iter_rows(max_row=1))) == cols

    def rows(name):
        it = wb[name].iter_rows(values_only=True)
        hdr = next(it)
        return [dict(zip(hdr, r)) for r in it]

    actions = rows("Actions")
    click = next(a for a in actions if a["Action"] == "button_clicked")
    assert click["Detail"] == "Build My University Plan" and click["Campaign"] == "Qya_First_BM"
    assert click["Attribution"] == "button_context" and click["Customer Phone/ID"] == "6281200000001"
    q = next(a for a in actions if a["Action"] == "question")
    assert q["Campaign"] == "Qya_First_BM"

    [camp] = rows("Campaigns")
    assert (camp["Recipients"], camp["Clicked Button"], camp["Responders"], camp["Questions"]) == (3, 1, 2, 1)

    recips = {r["Customer Phone/ID"]: r for r in rows("Broadcast Recipients")}
    assert recips["6281200000001"]["Clicked Button"] == "Build My University Plan"
    assert recips["6281200000002"]["Responded"] == "Yes" and recips["6281200000002"]["Clicked Button"] is None
    assert recips["6281299999999"]["Status"] == "failed"
    assert recips["6281299999999"]["Failure Reason"].startswith("Error 131026")

    leads = {r["Customer Phone/ID"]: r for r in rows("Leads")}
    assert leads["6281200000001"]["Lead Stage"] == "clicked"
    assert leads["6281200000002"]["Lead Stage"] == "question_received"
    assert leads["6281200000003"]["Lead Stage"] == "no_response"

    convs = rows("Conversations")
    msgs = rows("Messages")
    for c in convs:
        assert c["Total Messages"] == sum(1 for m in msgs if m["Room ID"] == c["Room ID"])


@respx.mock
def test_excel_writer_phone_column_is_text_format(settings, client, store, norm, tmp_path):
    _tenant(n_rooms=1)
    out = tmp_path / "o.xlsx"
    build(settings, client, store, norm).run(out_path=out)
    ws = load_workbook(out)["Messages"]
    cell = ws.cell(row=2, column=3)
    assert cell.value == "6281200000001" and cell.number_format == "@"


@respx.mock
def test_excel_writer_rows_ordered_by_room_then_timestamp(settings, client, store, norm, tmp_path):
    _tenant()
    out = tmp_path / "o.xlsx"
    build(settings, client, store, norm).run(out_path=out)
    rows = list(load_workbook(out, read_only=True)["Messages"].iter_rows(min_row=2, values_only=True))
    seen_rooms: list = []
    for r in rows:
        if not seen_rooms or seen_rooms[-1] != r[0]:
            assert r[0] not in seen_rooms
            seen_rooms.append(r[0])
    for room in seen_rooms:
        stamps = [r[5] for r in rows if r[0] == room]
        assert stamps == sorted(stamps)


@respx.mock
def test_excel_writer_rolls_over_at_excel_row_limit(settings, client, store, norm, tmp_path, monkeypatch):
    monkeypatch.setattr(excel_mod, "EXCEL_MAX_ROWS", 4)
    _tenant()
    out = tmp_path / "o.xlsx"
    build(settings, client, store, norm).run(out_path=out)
    wb = load_workbook(out, read_only=True)
    assert "Messages (2)" in wb.sheetnames
    assert next(wb["Messages (2)"].iter_rows(max_row=1, values_only=True))[0] == "Room ID"


@respx.mock
def test_formula_like_customer_text_is_neutralized(settings, client, store, norm, tmp_path):
    respx.post(ROOMS_URL).mock(return_value=httpx.Response(200, json=rooms_page([raw_room(1)], None)))
    respx.get(COMMENTS_URL).mock(return_value=httpx.Response(200, json=comments_page([
        comment(1, 491000001, email="6281200000001", text="=cmd|' /C calc'!A0")])))
    out = tmp_path / "o.xlsx"
    build(settings, client, store, norm, broadcasts=False).run(out_path=out)
    ws = load_workbook(out)["Messages"]
    cell = ws.cell(row=2, column=12)
    assert cell.data_type == "s" and cell.value.startswith("'=")
