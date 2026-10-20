"""Workbooks appear as soon as there is something to show, and say whether they are complete."""

import httpx
import pytest
import respx
from openpyxl import load_workbook

from qexport.load import excel as excel_mod
from qexport.load.excel import WorkbookLockedError, finish_locked_workbook
from qexport.orchestrator import EXIT_AUTH, EXIT_INTERRUPTED, EXIT_OK, NullProgress, RunCancelled
from tests.conftest import jobs_page, raw_job, raw_room, rooms_page
from tests.integration.test_pipeline import COMMENTS_URL, JOBS_URL, ROOMS_URL, _tenant, build


def _info(path) -> dict:
    ws = load_workbook(path, read_only=True)["Export Info"]
    return {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}


@respx.mock
def test_workbook_completeness_labeling(settings, client, store, norm):
    _tenant()
    result = build(settings, client, store, norm).run()
    assert result.exit_code == EXIT_OK
    final = result.report
    assert final.complete and not final.path.stem.endswith("_PARTIAL")
    assert _info(final.path)["Completeness"] == "Complete"


@respx.mock
def test_partial_workbook_names_what_is_missing(settings, client, store, norm):
    _tenant(forbidden=(3,))
    result = build(settings, client, store, norm).run()
    final = result.report
    assert not final.complete and final.path.stem.endswith("_PARTIAL")
    assert _info(final.path)["Completeness"] == "Partial (1 conversations failed)"


@respx.mock
def test_campaigns_not_sent_yet_do_not_make_a_workbook_partial(settings, client, store, norm):
    _tenant()
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page(
        [raw_job(900, "Qya_First_BM", recipients=3), raw_job(950, "Next week", status=4)], None)))
    result = build(settings, client, store, norm).run()
    assert result.report.complete


@respx.mock
def test_auto_workbook_intermediate_generation(settings, client, store, norm):
    _tenant()
    result = build(settings, client, store, norm).run()
    first, final = result.reports
    assert first.path != final.path and first.path.exists() and final.path.exists()
    assert first.rooms == 5 and first.broadcast_recipients == 0
    assert first.completeness == "Partial (1 broadcast campaigns not downloaded)"
    assert final.complete and final.broadcast_recipients == 3


@respx.mock
def test_no_intermediate_workbook_when_there_are_no_broadcasts_to_fetch(settings, client, store, norm):
    _tenant()
    respx.post(JOBS_URL).mock(return_value=httpx.Response(200, json=jobs_page([], None)))
    result = build(settings, client, store, norm).run()
    assert len(result.reports) == 1 and result.report.complete


@respx.mock
def test_auth_expiry_mid_run_still_writes_a_partial_workbook(settings, client, store, norm):
    _tenant(broken_auth=True)
    result = build(settings, client, store, norm).run()
    assert result.exit_code == EXIT_AUTH
    assert result.report is not None and result.report.path.exists()
    assert not result.report.complete and 0 < result.report.rooms


@respx.mock
def test_cancel_mid_run_writes_a_partial_workbook(settings, client, store, norm):
    class CancelOnce(NullProgress):
        n = 0

        def advance(self, n=1, **info):
            if "messages" in info:
                CancelOnce.n += 1
                if CancelOnce.n == 2:
                    raise RunCancelled
    _tenant()
    orch = build(settings, client, store, norm)
    orch._p = CancelOnce()
    result = orch.run()
    assert result.exit_code == EXIT_INTERRUPTED and "Cancelled" in result.message
    assert store.progress().done == 2
    assert result.report is not None and "3 conversations not downloaded" in result.report.completeness


@respx.mock
def test_no_workbook_when_nothing_was_downloaded(settings, client, store, norm):
    respx.post(ROOMS_URL).mock(return_value=httpx.Response(200, json=rooms_page([raw_room(1)], None)))
    respx.get(COMMENTS_URL).mock(return_value=httpx.Response(401))
    result = build(settings, client, store, norm, broadcasts=False).run()
    assert result.exit_code == EXIT_AUTH and result.reports == ()


@respx.mock
def test_excel_writer_locked_file_retry(settings, client, store, norm, tmp_path, monkeypatch):
    _tenant(n_rooms=1)
    out = tmp_path / "open_in_excel.xlsx"
    real_replace = excel_mod.os.replace
    locked = {"on": True}

    def replace(src, dst):
        if locked["on"] and str(dst) == str(out):
            raise PermissionError(13, "The process cannot access the file")
        return real_replace(src, dst)
    monkeypatch.setattr(excel_mod.os, "replace", replace)
    with pytest.raises(WorkbookLockedError) as ei:
        build(settings, client, store, norm, broadcasts=False).run(out_path=out)
    err = ei.value
    assert err.partial.exists() and not out.exists() and err.target == out

    with pytest.raises(WorkbookLockedError):          # still open
        finish_locked_workbook(err)
    locked["on"] = False                              # user closed Excel
    report = finish_locked_workbook(err)
    assert out.exists() and not err.partial.exists() and report.path == out
