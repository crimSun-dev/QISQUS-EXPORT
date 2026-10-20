"""Several tasks share one store file: writers queue, readers never wait, workbooks use snapshots."""

import sqlite3
import threading
import time

import pytest
from openpyxl import load_workbook

from qexport.api.models import RawComment, RawRoom
from qexport.load.excel import ExcelWriter
from qexport.orchestrator import build_workbook
from qexport.store.checkpoint import CheckpointStore
from tests.conftest import APP, comment, raw_room


def _seed(store, norm, n_rooms=3):
    rooms = [norm.room(RawRoom.model_validate(raw_room(i))) for i in range(1, n_rooms + 1)]
    store.upsert_rooms_page(rooms, None)
    store.mark_enumeration_done()
    store.mark_broadcast_enumeration_done()
    for r in rooms:
        rid = int(r.room_id)
        msgs = [norm.message(RawComment.model_validate(
            comment(rid * 10, rid, email=r.customer_identifier, text="Berapa biayanya?")), r)]
        store.commit_room(r.room_id, msgs, None)
    return rooms


def test_writer_waits_for_another_writer_instead_of_failing(store):
    other = sqlite3.connect(store.path, isolation_level=None, check_same_thread=False)
    other.execute("BEGIN IMMEDIATE")
    release = threading.Timer(0.5, lambda: other.execute("COMMIT"))
    release.start()
    t0 = time.monotonic()
    try:
        store.mark_run_completed()          # needs the write lock the other connection holds
    finally:
        release.join()
        other.close()
    assert time.monotonic() - t0 >= 0.4
    assert store.run_info()["last_completed_at_utc"] is not None


def test_reads_proceed_while_another_connection_is_writing(store, norm):
    _seed(store, norm, 1)
    other = sqlite3.connect(store.path, isolation_level=None)
    other.execute("BEGIN IMMEDIATE")
    other.execute("DELETE FROM messages")   # uncommitted
    t0 = time.monotonic()
    try:
        p = store.progress()
    finally:
        other.execute("ROLLBACK")
        other.close()
    assert time.monotonic() - t0 < 1.0 and p.messages == 1


def test_two_stores_on_one_file_commit_from_two_threads(tmp_path, norm):
    path = tmp_path / "shared.sqlite3"
    a, b = CheckpointStore(path, APP), CheckpointStore(path, APP)
    rooms = [norm.room(RawRoom.model_validate(raw_room(i))) for i in range(1, 41)]
    a.upsert_rooms_page(rooms, None)
    errors: list[Exception] = []

    def commit(store, chunk):
        try:
            for r in chunk:
                store.commit_room(r.room_id, [], None)
        except Exception as e:   # noqa: BLE001
            errors.append(e)
    threads = [threading.Thread(target=commit, args=(a, rooms[::2])),
               threading.Thread(target=commit, args=(b, rooms[1::2]))]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert errors == [] and a.progress().done == 40
    a.close()
    b.close()


def test_tasks_opening_a_brand_new_store_together_do_not_collide(tmp_path):
    path = tmp_path / "new.sqlite3"
    opened, errors = [], []

    def open_store():
        try:
            opened.append(CheckpointStore(path, APP))
        except Exception as e:   # noqa: BLE001
            errors.append(e)
    threads = [threading.Thread(target=open_store) for _ in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert errors == [] and len(opened) == 6
    for s in opened:
        s.close()


def test_snapshot_build_does_not_block_live_store(settings, store, norm, monkeypatch):
    rooms = _seed(store, norm)
    live_commit_s: list[float] = []
    real_write = ExcelWriter.write

    def write_while_downloading(self, snap, path, **kw):
        # A download commits to the live store while the workbook is being written.
        t0 = time.monotonic()
        store.commit_room(rooms[0].room_id, [], None)
        live_commit_s.append(time.monotonic() - t0)
        assert snap.path != store.path
        return real_write(self, snap, path, **kw)
    monkeypatch.setattr(ExcelWriter, "write", write_while_downloading)

    report = build_workbook(store, settings, ExcelWriter(norm))

    assert live_commit_s and live_commit_s[0] < 1.0
    assert store.counts()["actions"] == 0          # derived in the snapshot, live store untouched
    assert report.actions == 3
    assert list(store.path.parent.glob(".snapshot-*")) == []
    wb = load_workbook(report.path, read_only=True)
    assert len(list(wb["Actions"].iter_rows(min_row=2))) == 3


def test_snapshot_is_removed_when_the_build_fails(settings, store, norm, monkeypatch):
    _seed(store, norm, 1)

    def boom(self, *a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(ExcelWriter, "write", boom)
    with pytest.raises(RuntimeError):
        build_workbook(store, settings, ExcelWriter(norm))
    assert list(store.path.parent.glob(".snapshot-*")) == []


def test_snapshot_is_a_consistent_copy(store, norm, tmp_path):
    _seed(store, norm, 2)
    with store.snapshot() as snap:
        store.commit_room("491000001", [], None)       # live store changes after the snapshot
        assert snap.progress().messages == 2 and store.progress().messages == 1
        assert snap.integrity_check() == ["ok"]
