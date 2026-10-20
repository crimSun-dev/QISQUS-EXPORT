from __future__ import annotations

import os
import time
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from qexport import __version__
from qexport.store.checkpoint import GONE_PREFIX, CheckpointStore, Missing, from_iso
from qexport.transform.actions import (
    FORM_FIELD_LABELS, FORM_FIELDS, ActionType, lead_stage, merge_forms,
)
from qexport.transform.normalize import Normalizer, normalize_phone

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_CHARS = 32_767
_TRUNC_SUFFIX = "…[truncated]"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

LEADS_COLUMNS: tuple[str, ...] = (
    "Customer Phone/ID", "Customer Name", "Lead Stage", "Form Completion %",
    *(FORM_FIELD_LABELS[f] for f in FORM_FIELDS),
    "Buttons Clicked", "Button Clicks", "Questions", "Replies", "Total Actions",
    "First Action (WIB)", "Last Action (WIB)", "Last Attributed Campaign",
    "Broadcasts Received", "Best Broadcast Status", "Last Broadcast", "Last Broadcast Sent (WIB)",
    "Conversation Resolved", "Last Activity (WIB)", "Room ID",
)
ACTIONS_COLUMNS: tuple[str, ...] = (
    "Timestamp (WIB)", "Customer Phone/ID", "Customer Name", "Action", "Detail",
    "Campaign", "Attribution", "Form Fields", "Room ID", "Message ID", "Campaign Job ID",
)
CAMPAIGNS_COLUMNS: tuple[str, ...] = (
    "Broadcast Name", "Template", "Category", "Status", "Created (WIB)", "Started (WIB)",
    "Finished (WIB)", "Recipients", "Sent", "Delivered", "Read", "Failed",
    "Delivery Rate", "Read Rate", "Responders", "Clicked Button", "Click Rate",
    "Forms Submitted", "Questions", "Recipient Log", "Job ID",
)
RECIPIENTS_COLUMNS: tuple[str, ...] = (
    "Broadcast Name", "Template", "Sent At (WIB)", "Customer Phone/ID", "Customer Name",
    "Status", "Failure Reason", "Buttons Offered", "Clicked Button", "Responded",
    "Response Types", "First Response (WIB)", "Template Params", "Room ID", "WA Message ID", "Job ID",
)
CONVERSATIONS_COLUMNS: tuple[str, ...] = (
    "Room ID", "Customer Name", "Customer Phone/ID", "Channel", "Resolved", "Bot Handled",
    "First Message (WIB)", "Last Message (WIB)", "Total Messages", "Customer Msgs",
    "Agent Msgs", "Bot Msgs", "Broadcast Msgs", "Export State",
)
MESSAGES_COLUMNS: tuple[str, ...] = (
    "Room ID", "Customer Name", "Customer Phone/ID", "Channel", "Conversation Status",
    "Timestamp (WIB)", "Timestamp (UTC)", "Sender Type", "Sender Name", "Sender ID",
    "Message Type", "Message", "Button Clicked", "Broadcast Template", "Reply To",
    "Attachment Name", "Attachment URL", "Deleted", "Delivery Status", "Message ID",
)

_TEXT_COLUMNS = {"Customer Phone/ID", "Room ID", "Sender ID", "Phone Number", "Parent Number",
                 "WA Message ID", "Message ID", "Job ID", "Campaign Job ID"}
_PERCENT_COLUMNS = {"Form Completion %", "Delivery Rate", "Read Rate", "Click Rate"}
_WIDTHS = {"Message": 80, "Detail": 60, "Form Fields": 60, "Reply To": 40, "Attachment URL": 40,
           "Failure Reason": 40, "Template Params": 30, "Broadcast Name": 32, "Template": 26,
           "Campaign": 30, "Last Attributed Campaign": 30, "Last Broadcast": 30, "Buttons Clicked": 30,
           "Buttons Offered": 30, "Clicked Button": 26, "Interested Majors": 30, "WA Message ID": 30}
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F6F5C")
_DATE_FMT = "yyyy-mm-dd hh:mm:ss"


def sanitize_cell(value: str) -> tuple[str, bool]:
    """Strip XML-illegal control chars; neutralize formula/DDE injection by prefixing "'" when
    the value starts with = + - @ \\t \\r; truncate to EXCEL_MAX_CELL_CHARS with a suffix.
    Returns (clean_value, was_truncated)."""
    if value.startswith(_FORMULA_PREFIXES):
        value = "'" + value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > EXCEL_MAX_CELL_CHARS:
        return value[: EXCEL_MAX_CELL_CHARS - len(_TRUNC_SUFFIX)] + _TRUNC_SUFFIX, True
    return value, False


@dataclass(frozen=True)
class ExportReport:
    path: Path
    rooms: int
    messages: int
    truncated_cells: int
    skipped_rooms: list[str]
    sheets: list[str]
    duration_s: float
    actions: int = 0
    broadcast_jobs: int = 0
    broadcast_recipients: int = 0
    completeness: str = "Complete"      # Missing.describe(): "Complete" or "Partial (...)"

    @property
    def complete(self) -> bool:
        return self.completeness == "Complete"


class WorkbookLockedError(PermissionError):
    """The finished workbook could not replace `target` (open in Excel). It is kept, complete,
    at `partial`; finish_locked_workbook() retries the move once the file is closed."""

    def __init__(self, partial: Path, target: Path, report: ExportReport) -> None:
        super().__init__(f"{target.name} is open in another program (Excel?). Close it and retry; "
                         f"the new workbook is waiting at {partial.name}")
        self.partial, self.target, self.report = partial, target, report


def finish_locked_workbook(err: WorkbookLockedError) -> ExportReport:
    """Retry the final move. Raises WorkbookLockedError again while the target is still open."""
    try:
        os.replace(err.partial, err.target)
    except PermissionError as e:
        raise WorkbookLockedError(err.partial, err.target, err.report) from e
    return err.report


def expected_rows(store: CheckpointStore, include_recipients: bool = True) -> int:
    """Roughly how many rows write() will produce, for a progress bar."""
    c = store.counts()
    return (2 * c["rooms"] + c["actions"] + c["jobs"] + c["messages"]
            + (c["broadcast_logs"] if include_recipients else 0))


@dataclass
class _Counters:
    truncated: int = 0
    rows: dict[str, int] = field(default_factory=dict)


class _Sheet:
    """A write-only sheet with header styling, typed cells and rollover past Excel's row limit."""

    def __init__(self, wb: Workbook, title: str, headers: Sequence[str], counters: _Counters) -> None:
        self._wb, self._title, self._headers, self._c = wb, title, tuple(headers), counters
        self._part = 0
        self._ws = None
        self._rows = 0
        self.titles: list[str] = []
        self._kinds = ["text" if h in _TEXT_COLUMNS else "pct" if h in _PERCENT_COLUMNS else "auto"
                       for h in self._headers]
        self._open()

    def _open(self) -> None:
        self._close_current()
        self._part += 1
        title = self._title if self._part == 1 else f"{self._title} ({self._part})"
        ws = self._wb.create_sheet(title)
        for i, h in enumerate(self._headers, start=1):
            w = _WIDTHS.get(h, max(12, min(28, len(h) + 4)))
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        header = []
        for h in self._headers:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill = _HEADER_FONT, _HEADER_FILL
            c.alignment = Alignment(vertical="center")
            header.append(c)
        ws.append(header)
        self._ws, self._rows = ws, 1
        self.titles.append(title)

    def _close_current(self) -> None:
        if self._ws is not None and self._rows > 1:
            last = get_column_letter(len(self._headers))
            self._ws.auto_filter.ref = f"A1:{last}{self._rows}"

    def append(self, values: Sequence[object]) -> None:
        if self._rows >= EXCEL_MAX_ROWS:
            self._open()
        ws = self._ws
        row = []
        for kind, v in zip(self._kinds, values):
            if isinstance(v, str):
                v, cut = sanitize_cell(v)
                self._c.truncated += cut
            if kind == "text" and v is not None and not isinstance(v, str):
                v = str(v)
            if kind == "text" or kind == "pct" or isinstance(v, datetime):
                c = WriteOnlyCell(ws, value=v)
                c.number_format = "@" if kind == "text" else "0.0%" if kind == "pct" else _DATE_FMT
                row.append(c)
            else:
                row.append(v)
        ws.append(row)
        self._rows += 1
        self._c.rows[self._title] = self._c.rows.get(self._title, 0) + 1

    def finish(self) -> None:
        self._close_current()


class ExcelWriter:
    def __init__(self, normalizer: Normalizer) -> None:
        self._n = normalizer

    def _wib(self, dt: datetime | str | None) -> datetime | None:
        if isinstance(dt, str):
            dt = from_iso(dt)
        return self._n.to_local(dt) if dt is not None else None

    def write(self, store: CheckpointStore, out_path: Path, *,
              on_row: Callable[[int], None] | None = None,
              include_recipients: bool = True, missing: Missing | None = None) -> ExportReport:
        """Sheets: Leads, Actions, Campaigns, Broadcast Recipients, Conversations, Messages,
        Export Info. Writes to <out_path>.partial then os.replace() -> out_path (atomic).
        Raises WorkbookLockedError if out_path is open in Excel."""
        t0 = time.monotonic()
        missing = missing or store.missing()
        started = datetime.now(timezone.utc)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        partial = out_path.with_name(out_path.name + ".partial")
        wb = Workbook(write_only=True)
        cnt = _Counters()
        done = 0

        def tick() -> None:
            nonlocal done
            done += 1
            if on_row and done % 1000 == 0:
                on_row(done)

        rooms = {s.room.room_id: s for s in store.iter_room_summaries()}
        actions_by_room = store.actions_by_room()
        bc_by_phone = store.broadcast_status_by_phone()

        # ---- Leads
        leads = _Sheet(wb, "Leads", LEADS_COLUMNS, cnt)
        for room_id, s in rooms.items():
            r = s.room
            acts = actions_by_room.get(room_id, [])
            form = merge_forms(acts)
            bc = bc_by_phone.get(normalize_phone(r.customer_identifier), {})
            kinds = [a.action_type for a in acts]
            buttons = sorted({a.detail for a in acts if a.action_type == ActionType.BUTTON_CLICKED})
            attributed = [a.campaign_name for a in acts if a.campaign_name]
            leads.append([
                r.customer_identifier, form.get("name") or r.customer_name,
                lead_stage(acts, form, bc.get("best_status")) if s.fetch_state == "done" else "not_downloaded",
                len(form) / len(FORM_FIELDS) if form else 0.0,
                *(form.get(f) for f in FORM_FIELDS),
                " | ".join(buttons) or None, kinds.count(ActionType.BUTTON_CLICKED),
                kinds.count(ActionType.QUESTION), kinds.count(ActionType.REPLIED), len(acts),
                self._wib(acts[0].timestamp_utc) if acts else None,
                self._wib(acts[-1].timestamp_utc) if acts else None,
                attributed[-1] if attributed else None,
                bc.get("received", 0), bc.get("best_status"), bc.get("last_campaign"),
                self._wib(bc.get("last_sent_utc")),
                "Yes" if r.is_resolved else "No", self._wib(r.last_activity_utc), r.room_id,
            ])
            tick()
        leads.finish()

        # ---- Actions
        actions_sheet = _Sheet(wb, "Actions", ACTIONS_COLUMNS, cnt)
        n_actions = 0
        for room, a in store.iter_actions():
            form_txt = "; ".join(f"{FORM_FIELD_LABELS[k]}: {v}" for k, v in a.form_fields.items()) or None
            actions_sheet.append([
                self._wib(a.timestamp_utc), room.customer_identifier, room.customer_name,
                str(a.action_type), a.detail, a.campaign_name, str(a.attribution), form_txt,
                a.room_id, a.msg_id, a.campaign_job_id,
            ])
            n_actions += 1
            tick()
        actions_sheet.finish()

        # ---- Campaigns
        engagement = store.campaign_engagement()
        campaigns = _Sheet(wb, "Campaigns", CAMPAIGNS_COLUMNS, cnt)
        n_jobs = 0
        for j in store.iter_jobs():
            e = engagement.get(j["job_id"], {})
            total = j["total_recipient"] or 0
            reached = (j["total_delivered"] or 0) + (j["total_read"] or 0)
            campaigns.append([
                j["name"], j["template_name"], j["category"], j["status"],
                self._wib(j["created_at_utc"]), self._wib(j["started_at_utc"]), self._wib(j["finished_at_utc"]),
                total, j["total_sent"], j["total_delivered"], j["total_read"], j["total_failed"],
                reached / total if total else None, (j["total_read"] or 0) / total if total else None,
                e.get("responders", 0), e.get("clicked", 0),
                e.get("clicked", 0) / reached if reached else None,
                e.get("forms", 0), e.get("questions", 0),
                "not kept by Qiscus" if (j["last_error"] or "").startswith(GONE_PREFIX) else j["logs_state"],
                j["job_id"],
            ])
            n_jobs += 1
            tick()
        campaigns.finish()

        # ---- Broadcast Recipients
        n_recipients = 0
        recip_titles: list[str] = []
        if include_recipients:
            responses = store.recipient_responses()
            recipients = _Sheet(wb, "Broadcast Recipients", RECIPIENTS_COLUMNS, cnt)
            for lg in store.iter_broadcast_logs_for_export():
                resp = responses.get((lg["job_id"], lg["phone"]))
                recipients.append([
                    lg["job_name"], lg["template_name"], self._wib(lg["sent_at_utc"]), lg["phone"],
                    lg["customer_name"], lg["status"], lg["failure_reason"], lg["buttons"],
                    (resp[0] or None) if resp else None, "Yes" if resp else "No",
                    resp[1] if resp else None, self._wib(resp[2]) if resp else None,
                    lg["params"], lg["room_id"], lg["wamid"], lg["job_id"],
                ])
                n_recipients += 1
                tick()
            recipients.finish()
            recip_titles = recipients.titles

        # ---- Conversations
        convs = _Sheet(wb, "Conversations", CONVERSATIONS_COLUMNS, cnt)
        for s in rooms.values():
            r = s.room
            convs.append([
                r.room_id, r.customer_name, r.customer_identifier, r.source,
                "Yes" if r.is_resolved else "No", "Yes" if r.is_handled_by_bot else "No",
                self._wib(s.first_msg_utc), self._wib(s.last_msg_utc), s.message_count,
                s.customer_msgs, s.agent_msgs, s.bot_msgs, s.broadcast_msgs, s.fetch_state,
            ])
            tick()
        convs.finish()

        # ---- Messages
        msgs = _Sheet(wb, "Messages", MESSAGES_COLUMNS, cnt)
        n_msgs = 0
        for room, m in store.iter_messages_for_export():
            msgs.append([
                room.room_id, room.customer_name, room.customer_identifier, room.source,
                "Resolved" if room.is_resolved else "Open",
                self._wib(m.timestamp_utc), m.timestamp_utc.replace(tzinfo=None), str(m.sender_type),
                m.sender_name, m.sender_id, m.message_type, m.text, m.button_text, m.broadcast_template,
                m.reply_to_text, m.attachment_name, m.attachment_url, "Yes" if m.is_deleted else "No",
                m.delivery_status, m.msg_id,
            ])
            n_msgs += 1
            tick()
        msgs.finish()

        # ---- Export Info
        skipped = store.skipped()
        run = store.run_info()
        n_gone, newest_gone = store.jobs_gone()
        info = _Sheet(wb, "Export Info", ("Item", "Value"), cnt)
        rows: list[tuple[str, object]] = [
            ("Tool", f"qexport {__version__}"),
            ("Completeness", missing.describe()),
            ("App code", store.app_code),
            ("Run started (UTC)", run.get("started_at_utc")),
            ("Data as of (UTC)", run.get("last_completed_at_utc") or "incomplete run"),
            ("Workbook written (UTC)", started.isoformat(timespec="seconds")),
            ("Display timezone", str(self._n.tz)),
            ("Conversations", len(rooms)),
            ("Messages", n_msgs),
            ("Customer actions", n_actions),
            ("Broadcast jobs", n_jobs),
            ("Broadcast recipient rows", n_recipients if include_recipients else "not included"),
            ("Campaigns Qiscus no longer keeps recipients for",
             f"{n_gone:,} (created up to {newest_gone[:10]}; only their totals are available)"
             if n_gone else 0),
            ("Truncated cells", cnt.truncated),
            ("Skipped / failed items", len(skipped)),
            ("", ""),
            ("Definitions", ""),
            ("Lead Stage", "Furthest step reached: form_completed (>=60% of intake fields) > form_started > "
                           "clicked > question_received > replied > read > delivered > sent > failed > no_response"),
            ("Action: button_clicked", "Customer tapped a WhatsApp quick-reply button (e.g. 'Build My University Plan')"),
            ("Action: form_submitted", "Customer message with >= 3 intake-form fields (Name, Grade, Email, ...)"),
            ("Action: question", "Free text containing '?' or starting with a question word / 'mau tanya', 'mohon info'"),
            ("Action: replied", "Any other customer text"),
            ("Attribution: button_context", "Exact: the button tap references the broadcast's WhatsApp message id"),
            ("Attribution: reply_context", "Exact: the customer replied to (quoted) the broadcast message"),
            ("Attribution: last_touch", "Most recent broadcast to that phone within the attribution window"),
            ("Sent / Delivered / Read / Failed", "Exclusive buckets per recipient, as Qiscus reports them"),
            ("Delivery Rate", "(Delivered + Read) / Recipients"),
            ("Click Rate", "Distinct phones that clicked a button / (Delivered + Read)"),
            ("", ""),
            ("Privacy", "Contains personal data of students and parents (UU 27/2022 PDP). "
                        "Do not store on shared drives or forward without approval."),
        ]
        for k, v in rows:
            info.append([k, v])
        if skipped:
            info.append(["", ""])
            info.append(["Skipped item", "Reason"])
            for kind, ident, err in skipped:
                info.append([f"{kind} {ident}", err])
        info.finish()

        sheets = ["Leads", "Actions", "Campaigns", *recip_titles, "Conversations", *msgs.titles, "Export Info"]
        wb.save(partial)
        report = ExportReport(
            path=out_path, rooms=len(rooms), messages=n_msgs, truncated_cells=cnt.truncated,
            skipped_rooms=[f"{k} {i}: {e}" for k, i, e in skipped], sheets=sheets,
            duration_s=time.monotonic() - t0, actions=n_actions, broadcast_jobs=n_jobs,
            broadcast_recipients=n_recipients, completeness=missing.describe())
        try:
            os.replace(partial, out_path)
        except PermissionError as e:    # target open in Excel; the finished .partial is kept
            raise WorkbookLockedError(partial, out_path, report) from e
        return report
